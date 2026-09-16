"""Exportação de relatórios em Excel (.xlsx), no estilo painel de BI:
KPIs, abas por dimensão e gráficos — não é um CSV cru."""
import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ACCENT = "1E6FD9"
_ACCENT_DEEP = "1857AD"
_INK = "0F2942"
_MUTED = "64748B"
_BORDER = Border(*([Side(style="thin", color="D9E0EA")] * 4))

_RISCO_CORES = {
    "VERMELHO": ("FEE2E2", "991B1B"),
    "LARANJA": ("FFEDD5", "9A3412"),
    "AMARELO": ("FEF9C3", "854D0E"),
    "VERDE": ("DCFCE7", "166534"),
    "AZUL": ("DBEAFE", "1E40AF"),
}


def _titulo_sheet(ws, titulo: str, subtitulo: str = "") -> None:
    ws["A1"] = titulo
    ws["A1"].font = Font(size=16, bold=True, color=_INK)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(size=10, color=_MUTED)


def _tabela(ws, linha_inicial: int, cabecalho: list[str], linhas: list[tuple],
            larguras: list[int] | None = None):
    """Escreve uma tabela com cabeçalho estilizado e devolve a linha seguinte."""
    for j, titulo in enumerate(cabecalho, start=1):
        cel = ws.cell(row=linha_inicial, column=j, value=titulo)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=_ACCENT)
        cel.alignment = Alignment(horizontal="left", vertical="center")
        cel.border = _BORDER
    for i, linha in enumerate(linhas, start=linha_inicial + 1):
        for j, valor in enumerate(linha, start=1):
            cel = ws.cell(row=i, column=j, value=valor)
            cel.border = _BORDER
            if i % 2 == 0:
                cel.fill = PatternFill("solid", fgColor="F4F7FB")
    if larguras:
        for j, largura in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(j)].width = largura
    ws.freeze_panes = ws.cell(row=linha_inicial + 1, column=1)
    return linha_inicial + len(linhas) + 1


def _grafico_barras(ws, ancora: str, titulo: str, dados_ref, categorias_ref):
    chart = BarChart()
    chart.title = titulo
    chart.style = 10
    chart.y_axis.title = "Atendimentos"
    chart.x_axis.title = None
    chart.height, chart.width = 7, 13
    chart.add_data(dados_ref, titles_from_data=True)
    chart.set_categories(categorias_ref)
    chart.legend = None
    ws.add_chart(chart, ancora)


def gerar_producao_xlsx(dados: dict) -> bytes:
    periodo = f"{dados['periodo_de'].strftime('%d/%m/%Y')} a {dados['periodo_ate'].strftime('%d/%m/%Y')}"
    wb = Workbook()

    # ── Painel ──────────────────────────────────────────────
    painel = wb.active
    painel.title = "Painel"
    _titulo_sheet(painel, "Relatório de Produção — Vida+ Clínica", f"Período: {periodo}")

    kpis = [
        ("Atendimentos finalizados", dados["total"]),
        ("Profissionais com produção", len(dados["por_profissional"])),
        ("Diagnósticos distintos (CID/CIAP)", len(dados["por_cid"])),
    ]
    for i, (rotulo, valor) in enumerate(kpis):
        col = 1 + i * 2
        c1 = painel.cell(row=4, column=col, value=rotulo)
        c1.font = Font(size=9, color=_MUTED)
        c2 = painel.cell(row=5, column=col, value=valor)
        c2.font = Font(size=20, bold=True, color=_ACCENT_DEEP)
    painel.column_dimensions["A"].width = 26
    painel.column_dimensions["C"].width = 26
    painel.column_dimensions["E"].width = 26

    # tabelas de apoio (fora da área visível dos KPIs) que alimentam os gráficos
    linha = 8
    painel.cell(row=linha, column=1, value="Por profissional").font = Font(bold=True, color=_INK)
    linha += 1
    prof_ini = linha
    linha = _tabela(painel, linha, ["Profissional", "Total"],
                    [(r["nome"], r["total"]) for r in dados["por_profissional"]], [26, 10])
    if dados["por_profissional"]:
        _grafico_barras(
            painel, f"D{prof_ini}", "Atendimentos por profissional",
            Reference(painel, min_col=2, min_row=prof_ini, max_row=linha - 1),
            Reference(painel, min_col=1, min_row=prof_ini + 1, max_row=linha - 1),
        )

    linha += 16
    painel.cell(row=linha, column=1, value="Por desfecho").font = Font(bold=True, color=_INK)
    linha += 1
    des_ini = linha
    des_linhas = [(r["nome"].replace("_", " ").title(), r["total"]) for r in dados["por_desfecho"]]
    linha = _tabela(painel, linha, ["Desfecho", "Total"], des_linhas, [26, 10])
    if des_linhas:
        _grafico_barras(
            painel, f"D{des_ini}", "Atendimentos por desfecho",
            Reference(painel, min_col=2, min_row=des_ini, max_row=linha - 1),
            Reference(painel, min_col=1, min_row=des_ini + 1, max_row=linha - 1),
        )

    # ── Por profissional (tabela completa) ─────────────────────
    ws = wb.create_sheet("Por profissional")
    _titulo_sheet(ws, "Produção por profissional", periodo)
    _tabela(ws, 4, ["Profissional", "Atendimentos"],
            [(r["nome"], r["total"]) for r in dados["por_profissional"]], [30, 14])

    # ── Por desfecho ────────────────────────────────────────
    ws = wb.create_sheet("Por desfecho")
    _titulo_sheet(ws, "Produção por desfecho", periodo)
    _tabela(ws, 4, ["Desfecho", "Atendimentos"], des_linhas, [30, 14])

    # ── Por risco (com cor por classificação) ──────────────
    ws = wb.create_sheet("Por risco")
    _titulo_sheet(ws, "Atendimentos por classificação de risco", periodo)
    linha_ini = 4
    linha_fim = _tabela(ws, linha_ini, ["Risco", "Atendimentos"],
                        [(r["nome"], r["total"]) for r in dados["por_risco"]], [22, 14])
    for i, r in enumerate(dados["por_risco"], start=linha_ini + 1):
        cores = _RISCO_CORES.get(r["nome"])
        if cores:
            bg, fg = cores
            cel = ws.cell(row=i, column=1)
            cel.fill = PatternFill("solid", fgColor=bg)
            cel.font = Font(bold=True, color=fg)

    # ── CID/CIAP ────────────────────────────────────────────
    ws = wb.create_sheet("CID e CIAP")
    _titulo_sheet(ws, "Diagnósticos mais frequentes (CID-10 / CIAP-2)", periodo)
    _tabela(ws, 4, ["Código", "Descrição", "Ocorrências"],
            [(r["codigo"], r["descricao"], r["total"]) for r in dados["por_cid"]], [12, 50, 14])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
